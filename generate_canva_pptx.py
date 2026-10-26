import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    bg_img_path = "dubai_emerald_skyline_bg.png"
    if not os.path.exists(bg_img_path):
        import create_background
        create_background.generate_dubai_background()

    # Color Tokens from Canva reference
    BG_EMERALD = RGBColor(11, 46, 33)       # #0B2E21
    CARD_BG = RGBColor(9, 38, 27)           # #09261B deep container
    CARD_BG_LIGHT = RGBColor(14, 56, 40)    # #0E3828
    CARD_BORDER = RGBColor(40, 110, 80)     # #286E50
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MINT = RGBColor(167, 243, 208)     # #A7F3D0
    TEXT_BODY = RGBColor(226, 245, 236)     # #E2F5EC
    TEXT_MUTED = RGBColor(140, 195, 170)    # #8CC3AA
    ACCENT_GOLD = RGBColor(253, 230, 138)   # #FDE68A

    FONT_TITLE = "Montserrat"
    FONT_BODY = "Arial"

    def apply_slide_bg(slide):
        # Add full-bleed background image of Dubai wireframe skyline on emerald green
        slide.shapes.add_picture(bg_img_path, 0, 0, Inches(13.333), Inches(7.5))

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_header(slide, section_tag, title_text):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.95))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = section_tag.upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT

        p2 = tf.add_paragraph()
        p2.text = title_text.upper()
        p2.font.name = FONT_TITLE
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(3)

    # ==============================================================
    # SLIDE 1: COVER (Exactly matching user Canva screenshot style)
    # ==============================================================
    s1 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s1)

    tx_box1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.333), Inches(3.8))
    tf1 = tx_box1.text_frame
    tf1.word_wrap = True
    tf1.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf1.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "The Youth Financial Wellbeing Dubai"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Every Dirham Has A Plan"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(46)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(8)

    p3 = tf1.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "COM345 Communication Plan"
    p3.font.name = FONT_TITLE
    p3.font.size = Pt(20)
    p3.font.color.rgb = TEXT_MINT
    p3.space_before = Pt(14)

    p4 = tf1.add_paragraph()
    p4.alignment = PP_ALIGN.CENTER
    p4.text = 'Campaign Message: “Dubai Life. Smarter Finances.”'
    p4.font.name = FONT_BODY
    p4.font.size = Pt(14)
    p4.font.color.rgb = TEXT_BODY
    p4.space_before = Pt(6)

    # Team members badge card at bottom
    add_card(s1, Inches(2.5), Inches(5.8), Inches(8.333), Inches(0.85), bg_color=CARD_BG, border_color=CARD_BORDER)
    tx_t = s1.shapes.add_textbox(Inches(2.6), Inches(5.95), Inches(8.133), Inches(0.6))
    tf_t = tx_t.text_frame
    p_t = tf_t.paragraphs[0]
    p_t.alignment = PP_ALIGN.CENTER
    p_t.text = "PROJECT TEAM: Lareb's Part  •  Gleb's Part  •  Amaliya's Part  •  Akbar's Part"
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_MINT

    # ==============================================================
    # SLIDE 2: Lareb’s Part — 1. Background (Macro Factors)
    # ==============================================================
    s2 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s2)
    add_header(s2, "Lareb’s Part • 1. Background", "Rising Cost-of-Living Pressures in Dubai")

    # 4 Cards for the 4 data points in Background
    items_s2 = [
        ("Housing as Main Driver", "Dubai's population is growing faster than new homes are built, with around 470 new residents arriving daily (The National, 2025)."),
        ("Rents vs. Salaries", "Rents are expected to rise by up to 6% in 2026 (Abbas, 2025), while average salaries were projected to see no increase in 2025 (ValuStrat, 2025)."),
        ("Transport Cost Growth", "Transport costs have also grown, with peak-hour Salik tolls rising from AED 4 to AED 6 (Gulf Business, 2024)."),
        ("Buy Now, Pay Later (BNPL)", "Buy now, pay later data from Tabby and Tamara now appears on UAE credit reports, so everyday instalment purchases can affect future loans and rentals (Gulf News, 2026).")
    ]

    # Intro banner
    add_card(s2, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.8), bg_color=CARD_BG)
    tb_intro2 = s2.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.333), Inches(0.6))
    tf_i2 = tb_intro2.text_frame
    tf_i2.word_wrap = True
    p = tf_i2.paragraphs[0]
    p.text = "Young people in Dubai aged 25 to 34 face rising cost-of-living pressures."
    p.font.name = FONT_TITLE
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    w_card = Inches(5.7)
    h_card = Inches(2.05)
    for i, (title, desc) in enumerate(items_s2):
        cx = Inches(0.8) + (i % 2) * Inches(6.033)
        cy = Inches(2.45) + (i // 2) * Inches(2.25)
        add_card(s2, cx, cy, w_card, h_card, bg_color=CARD_BG)
        
        tb = s2.shapes.add_textbox(cx + Inches(0.25), cy + Inches(0.2), w_card - Inches(0.5), h_card - Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(6)

    # ==============================================================
    # SLIDE 3: Lareb’s Part — 1. Background (Communication Issue)
    # ==============================================================
    s3 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s3)
    add_header(s3, "Lareb’s Part • 1. Background", "The Communication Issue")

    add_card(s3, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2), bg_color=CARD_BG)
    tb_s3 = s3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.933), Inches(4.6))
    tf_s3 = tb_s3.text_frame
    tf_s3.word_wrap = True

    p = tf_s3.paragraphs[0]
    p.text = "EXISTING SUPPORT VS. UNAWARENESS"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_s3.add_paragraph()
    p2.text = (
        "The main communication issue is that support already exists, such as the "
        "Dubai Land Department rental index and Etihad Credit Bureau credit reports, "
        "but many young people are unaware of it."
    )
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(10)

    p3 = tf_s3.add_paragraph()
    p3.text = "RELIANCE ON SOCIAL MEDIA & REGULATORY RESPONSE"
    p3.font.name = FONT_TITLE
    p3.font.size = Pt(15)
    p3.font.bold = True
    p3.font.color.rgb = TEXT_MINT
    p3.space_before = Pt(24)

    p4 = tf_s3.add_paragraph()
    p4.text = (
        "Instead, they rely on social media, where unrealistic lifestyles and unverified advice are common, "
        "a risk the UAE has addressed by licensing finance influencers (Issa, 2025)."
    )
    p4.font.name = FONT_BODY
    p4.font.size = Pt(13)
    p4.font.color.rgb = TEXT_WHITE
    p4.space_before = Pt(10)

    # ==============================================================
    # SLIDE 4: Lareb’s Part — 2. Problems and Opportunities
    # ==============================================================
    s4 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s4)
    add_header(s4, "Lareb’s Part • 2. Problems and Opportunities", "Problem and Opportunity Analysis")

    # Left: Problem
    add_card(s4, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2), bg_color=CARD_BG)
    tb_pr = s4.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.6))
    tf_pr = tb_pr.text_frame
    tf_pr.word_wrap = True

    p = tf_pr.paragraphs[0]
    p.text = "THE PROBLEM"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_pr.add_paragraph()
    p2.text = (
        "The problem is that young people in Dubai are facing rising living costs but lack clear and reliable "
        "guidance on how to manage them.\n\n"
        "Support services and educational resources already exist, but young people are unaware of them, "
        "and the active social media life spreads unrealistic expectations and misinformation about financial well-being."
    )
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_BODY
    p2.space_before = Pt(12)

    # Right: Opportunity
    add_card(s4, Inches(6.833), Inches(1.5), Inches(5.7), Inches(5.2), bg_color=CARD_BG)
    tb_op = s4.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.1), Inches(4.6))
    tf_op = tb_op.text_frame
    tf_op.word_wrap = True

    p = tf_op.paragraphs[0]
    p.text = "THE OPPORTUNITY"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_op.add_paragraph()
    p2.text = (
        "This is an opportunity for YFWD to step in early, before it becomes a broader social problem.\n\n"
        "As a neutral organisation, YFWD can position itself as a trusted source of guidance by using "
        "platforms young people already use to promote financial literacy, responsible spending and saving, "
        "and awareness of available support."
    )
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_BODY
    p2.space_before = Pt(12)

    # ==============================================================
    # SLIDE 5: Gleb’s Part — Target Publics & Message
    # ==============================================================
    s5 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s5)
    add_header(s5, "Gleb’s Part", "Target Publics & Core Message")

    # Target Publics Card
    add_card(s5, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2), bg_color=CARD_BG)
    tb_tp = s5.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.6))
    tf_tp = tb_tp.text_frame
    tf_tp.word_wrap = True

    p = tf_tp.paragraphs[0]
    p.text = "TARGET PUBLICS"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_tp.add_paragraph()
    p2.text = "Primary Public:"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(14)

    p3 = tf_tp.add_paragraph()
    p3.text = "Dubai residents aged 25-34, especially young professionals, graduates and renters."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12)
    p3.font.color.rgb = TEXT_BODY
    p3.space_before = Pt(4)

    p4 = tf_tp.add_paragraph()
    p4.text = "Secondary Publics:"
    p4.font.name = FONT_TITLE
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = TEXT_WHITE
    p4.space_before = Pt(20)

    p5 = tf_tp.add_paragraph()
    p5.text = "Universities, employers, banks, property providers, discount platforms and community organisations."
    p5.font.name = FONT_BODY
    p5.font.size = Pt(12)
    p5.font.color.rgb = TEXT_BODY
    p5.space_before = Pt(4)

    # Message Card
    add_card(s5, Inches(6.833), Inches(1.5), Inches(5.7), Inches(5.2), bg_color=CARD_BG)
    tb_msg = s5.shapes.add_textbox(Inches(7.133), Inches(1.8), Inches(5.1), Inches(4.6))
    tf_msg = tb_msg.text_frame
    tf_msg.word_wrap = True

    p = tf_msg.paragraphs[0]
    p.text = "MESSAGE"
    p.font.name = FONT_TITLE
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_msg.add_paragraph()
    p2.text = '“Dubai Life. Smarter Finances.”'
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(24)

    p3 = tf_msg.add_paragraph()
    p3.text = "Helping young people manage Dubai’s cost of living through informed financial decisions, practical planning and access to real cost-saving opportunities."
    p3.font.name = FONT_BODY
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_BODY
    p3.space_before = Pt(14)

    # ==============================================================
    # SLIDE 6: Gleb’s Part — Strategy: Educate - Equip - Connect
    # ==============================================================
    s6 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s6)
    add_header(s6, "Gleb’s Part • Strategy", "Educate - Equip - Connect")

    s_items = [
        ("Educate", "Social media content, workshops and practical advice on budgeting, rent, transport and saving."),
        ("Equip", "Provide budget-planning software, financial tools and cost-saving guides."),
        ("Connect", "Partner with banks, rental/property providers, Fazaa and other discount platforms, restaurants and food-delivery services to provide relevant financial benefits and discounts.\n\n• Zywad bank card: Seek educational collaboration with Zywa, using its money-management app and prepaid-card model to explain spending control.")
    ]

    w_col = Inches(3.72)
    gap_col = Inches(0.28)
    for i, (title, desc) in enumerate(s_items):
        cx = Inches(0.8) + i * (w_col + gap_col)
        add_card(s6, cx, Inches(1.5), w_col, Inches(5.2), bg_color=CARD_BG)
        
        tb = s6.shapes.add_textbox(cx + Inches(0.25), Inches(1.8), w_col - Inches(0.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"0{i+1}. {title.upper()}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(12)

    # ==============================================================
    # SLIDE 7: Gleb’s Part — Savings Example Calculation
    # ==============================================================
    s7 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s7)
    add_header(s7, "Gleb’s Part", "Savings Example Calculation")

    add_card(s7, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2), bg_color=CARD_BG)
    tb_s7 = s7.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.933), Inches(4.6))
    tf_s7 = tb_s7.text_frame
    tf_s7.word_wrap = True

    p = tf_s7.paragraphs[0]
    p.text = "HOW THE SAME BUDGET CAN BE USED BETTER THROUGH PROPOSED SAVINGS"
    p.font.name = FONT_TITLE
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    p2 = tf_s7.add_paragraph()
    p2.text = (
        "Meal planning, suitable transport alternatives, fewer food-delivery orders and cancelling "
        "unused subscriptions could save AED 460 monthly."
    )
    p2.font.name = FONT_BODY
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_WHITE
    p2.space_before = Pt(12)

    # Big numbers highlight
    p3 = tf_s7.add_paragraph()
    p3.text = "Within the same AED 6,000 budget, savings could increase from AED 500 to AED 960."
    p3.font.name = FONT_TITLE
    p3.font.size = Pt(20)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_GOLD
    p3.space_before = Pt(20)

    p4 = tf_s7.add_paragraph()
    p4.text = "This is an illustrative example; savings are not guaranteed or a card benefit."
    p4.font.name = FONT_BODY
    p4.font.size = Pt(11.5)
    p4.font.italic = True
    p4.font.color.rgb = TEXT_MUTED
    p4.space_before = Pt(16)

    # ==============================================================
    # SLIDE 8: Amaliya’s Part — Channels and Tactics
    # ==============================================================
    s8 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s8)
    add_header(s8, "Amaliya’s Part • Channels and Tactics", "Educate - Equip - Connect Implementation")

    tactics = [
        ("Educate:", "• Publish two weekly Instagram and TikTok posts about budgeting, rent, transport and saving.\n\n• Run six practical workshops through universities, employers and community organisations."),
        ("Equip:", "• Create a campaign webpage with a free budget planner, savings calculator and practical guides.\n\n• Promote it through QR codes and partner newsletters.\n\n• Launch a \"30-Day Budget Challenge\" to help participants track spending and set savings goals."),
        ("Connect:", "• Seek educational collaboration with Zywa, using its money-management app and prepaid-card model to explain spending control.\n\n• Seek partnerships with banks, property providers, restaurants and discount platforms, including Fazaa, to share cost-saving opportunities.\n\n• Clearly explain eligibility and offer conditions.")
    ]

    for i, (title, desc) in enumerate(tactics):
        cx = Inches(0.8) + i * (w_col + gap_col)
        add_card(s8, cx, Inches(1.5), w_col, Inches(5.2), bg_color=CARD_BG)
        
        tb = s8.shapes.add_textbox(cx + Inches(0.25), Inches(1.8), w_col - Inches(0.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.name = FONT_TITLE
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(10)

    # ==============================================================
    # SLIDE 9: Amaliya’s Part — Budget (12 Months, Total AED 50,000)
    # ==============================================================
    s9 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s9)
    add_header(s9, "Amaliya’s Part • Budget", "Duration: 12 Months. Total: AED 50,000")

    # Left: Table
    add_card(s9, Inches(0.8), Inches(1.5), Inches(6.5), Inches(5.2), bg_color=CARD_BG)
    rows_b, cols_b = 11, 2
    table_shape = s9.shapes.add_table(rows_b, cols_b, Inches(1.0), Inches(1.7), Inches(6.1), Inches(4.7))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(4.3)
    tbl.columns[1].width = Inches(1.8)

    b_data = [
        ["Research and evaluation", "5,000"],
        ["Educational content", "9,000"],
        ["Social media advertising", "10,000"],
        ["Campaign webpage tools", "6,000"],
        ["Six workshops", "9,000"],
        ["Creator collaborations", "4,000"],
        ["Printing and QR posters", "2,000"],
        ["Partner coordination", "2,000"],
        ["Contingency", "3,000"],
        ["Total", "50,000"]
    ]

    cell_h0 = tbl.cell(0, 0)
    cell_h0.text = "Item"
    cell_h1 = tbl.cell(0, 1)
    cell_h1.text = "AED"
    for c in [cell_h0, cell_h1]:
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG_LIGHT
        for p in c.text_frame.paragraphs:
            p.font.name = FONT_TITLE
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = TEXT_MINT

    for r, (item, aed) in enumerate(b_data):
        c0 = tbl.cell(r + 1, 0)
        c0.text = item
        c1 = tbl.cell(r + 1, 1)
        c1.text = aed
        for c in [c0, c1]:
            c.fill.solid()
            c.fill.fore_color.rgb = CARD_BG
            for p in c.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(9.5)
                if r == 9: # Total row
                    p.font.bold = True
                    p.font.color.rgb = ACCENT_GOLD
                else:
                    p.font.color.rgb = TEXT_BODY

    # Right: Cost Controls Notes
    add_card(s9, Inches(7.5), Inches(1.5), Inches(5.033), Inches(5.2), bg_color=CARD_BG)
    tb_c = s9.shapes.add_textbox(Inches(7.8), Inches(1.8), Inches(4.433), Inches(4.6))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "COST CONTROL GUIDELINES"
    p.font.name = FONT_TITLE
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT

    guidelines = [
        "Seek free venues, volunteer speakers and promotion through partner networks.",
        "Use existing YFWD staff and reusable materials to control costs.",
        "Confirm supplier prices before spending and review the budget monthly."
    ]
    for g in guidelines:
        p_g = tf_c.add_paragraph()
        p_g.text = f"• {g}"
        p_g.font.name = FONT_BODY
        p_g.font.size = Pt(11)
        p_g.font.color.rgb = TEXT_BODY
        p_g.space_before = Pt(14)

    # ==============================================================
    # SLIDE 10: Akbar’s part — 3. Objectives (SMART)
    # ==============================================================
    s10 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s10)
    add_header(s10, "Akbar’s Part • 3. Objectives", "SMART Campaign Objectives")

    # Intro
    add_card(s10, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.65), bg_color=CARD_BG)
    tb_i10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.333), Inches(0.5))
    tf_i10 = tb_i10.text_frame
    p = tf_i10.paragraphs[0]
    p.text = "To support the Educate, Equip, Connect strategy, the campaign establishes three SMART objectives:"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_BODY

    objs = [
        ("Objective 1 (Awareness / Educate):", "To increase awareness of financial literacy and YFWD support resources among Dubai residents aged 25–34 by achieving 150,000 total impressions and a 4% average engagement rate on Instagram and TikTok campaign content within 6 months of launch."),
        ("Objective 2 (Behavioral Change / Equip):", "To drive 5,000 downloads or active uses of the campaign webpage’s free budget planner and guides, and to enroll 1,000 young professionals in the \"30-Day Budget Challenge\" within 12 months."),
        ("Objective 3 (Partnership & Engagement / Connect):", "To engage at least 300 total participants across six interactive workshops and establish at least 4 active cost-saving partnerships (e.g., Zywa, Fazaa, rental/lifestyle platforms) by the end of the 12-month campaign period.")
    ]

    for i, (title, desc) in enumerate(objs):
        cx = Inches(0.8) + i * (w_col + gap_col)
        add_card(s10, cx, Inches(2.25), w_col, Inches(4.45), bg_color=CARD_BG)
        
        tb = s10.shapes.add_textbox(cx + Inches(0.25), Inches(2.45), w_col - Inches(0.5), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(10)

    # ==============================================================
    # SLIDE 11: Akbar’s part — 10. Evaluation (Beginning: Formative)
    # ==============================================================
    s11 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s11)
    add_header(s11, "Akbar’s Part • 10. Evaluation and Modifications", "Beginning: Baseline Evaluation (Formative)")

    # Intro
    add_card(s11, Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.85), bg_color=CARD_BG)
    tb_i11 = s11.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.333), Inches(0.65))
    tf_i11 = tb_i11.text_frame
    tf_i11.word_wrap = True
    p = tf_i11.paragraphs[0]
    p.text = "Evaluation is conducted continuously across three distinct phases: Beginning (Formative Research), Middle (Process Evaluation & Modification), and End (Summative Evaluation), to ensure the AED 50,000 budget delivers maximum impact and meets all campaign objectives."
    p.font.name = FONT_BODY
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_BODY

    # 2 Cards for Beginning
    w_half = Inches(5.7)
    add_card(s11, Inches(0.8), Inches(2.5), w_half, Inches(4.2), bg_color=CARD_BG)
    tb_b1 = s11.shapes.add_textbox(Inches(1.1), Inches(2.8), Inches(5.1), Inches(3.6))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "Desktop Research & Social Listening:"
    p.font.name = FONT_TITLE
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT
    p2 = tf_b1.add_paragraph()
    p2.text = "Monitor online discussions, hashtags, and social media trends surrounding Dubai living costs and financial management among young professionals (aged 25–34) before launch."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_BODY
    p2.space_before = Pt(12)

    add_card(s11, Inches(6.833), Inches(2.5), w_half, Inches(4.2), bg_color=CARD_BG)
    tb_b2 = s11.shapes.add_textbox(Inches(7.133), Inches(2.8), Inches(5.1), Inches(3.6))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "Pre-Campaign Survey:"
    p.font.name = FONT_TITLE
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_MINT
    p2 = tf_b2.add_paragraph()
    p2.text = "Conduct a short digital baseline survey (distributed through social media and university/employer channels) to measure existing awareness levels of YFWD resources, current spending habits, and top financial stress points."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_BODY
    p2.space_before = Pt(12)

    # ==============================================================
    # SLIDE 12: Akbar’s part — 10. Evaluation (Middle: Process)
    # ==============================================================
    s12 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s12)
    add_header(s12, "Akbar’s Part • 10. Evaluation and Modifications", "Middle: Ongoing Monitoring & Modification (Process Evaluation)")

    mid_items = [
        ("Monthly Social Media Analytics:", "Track reach, engagement, click-through rates, and follower growth on Instagram and TikTok. If engagement falls below the target 4%, pivot content formats toward higher-performing Reels, creator collaborations, or interactive polls."),
        ("Webpage & Tool Tracking:", "Use Google Analytics to track visits to the campaign webpage, QR code scans from posters, downloads of the budget planner tool, and completion rates for the \"30-Day Budget Challenge.\""),
        ("Mid-Campaign Review (Month 6):", "Assess workshop attendance and feedback forms. Reallocate unused portions of the AED 3,000 contingency budget toward higher-performing tactics (e.g., boosting social media advertising if web traffic needs a push).")
    ]

    for i, (title, desc) in enumerate(mid_items):
        cx = Inches(0.8) + i * (w_col + gap_col)
        add_card(s12, cx, Inches(1.5), w_col, Inches(5.2), bg_color=CARD_BG)
        
        tb = s12.shapes.add_textbox(cx + Inches(0.25), Inches(1.8), w_col - Inches(0.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(10)

    # ==============================================================
    # SLIDE 13: Akbar’s part — 10. Evaluation (End: Summative)
    # ==============================================================
    s13 = prs.slides.add_slide(blank_layout)
    apply_slide_bg(s13)
    add_header(s13, "Akbar’s Part • 10. Evaluation and Modifications", "End: Final Measurement (Summative Evaluation)")

    end_items = [
        ("Post-Campaign Survey:", "Re-run the digital survey at Month 12 to measure shifts in financial awareness, self-reported saving behavior, and perception of YFWD as a trusted financial guidance source."),
        ("SMART Objective Audit:", "• Awareness: Total social media impressions vs. target (150,000).\n\n• Equip/Behavior: Total budget planner downloads (Target: 5,000) and 30-Day Challenge participants (Target: 1,000).\n\n• Connect: Total workshop attendance (Target: 300) and signed/active partner offers (Target: 4)."),
        ("ROI & Future Planning:", "Summarize campaign reach per AED spent and compile a final evaluation report with recommendations for YFWD's long-term financial literacy initiatives.")
    ]

    for i, (title, desc) in enumerate(end_items):
        cx = Inches(0.8) + i * (w_col + gap_col)
        add_card(s13, cx, Inches(1.5), w_col, Inches(5.2), bg_color=CARD_BG)
        
        tb = s13.shapes.add_textbox(cx + Inches(0.25), Inches(1.8), w_col - Inches(0.5), Inches(4.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_TITLE
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_MINT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_BODY
        p2.space_before = Pt(10)

    out_file = "Dubai_Life_Smarter_Finances_COM345.pptx"
    prs.save(out_file)
    print(f"Successfully generated {out_file} with {len(prs.slides)} slides matching document order and Canva emerald background!")

if __name__ == "__main__":
    build_presentation()
