import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen layout (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Signature DFM Arena Color Palette
    COBALT = RGBColor(0, 82, 255)        # #0052FF Signature Electric Blue
    COBALT_DARK = RGBColor(0, 50, 160)   # #0032A0
    MIDNIGHT = RGBColor(7, 10, 18)       # #070A12 Deep Midnight Background
    MIDNIGHT_CARD = RGBColor(14, 21, 38) # #0E1526
    MIDNIGHT_CARD2 = RGBColor(19, 29, 52)# #131D34
    WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    TEXT_BODY = RGBColor(203, 213, 225)  # #CBD5E1
    ROSE = RGBColor(239, 68, 68)         # #EF4444
    GOLD = RGBColor(245, 158, 11)        # #F59E0B

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_card(slide, left, top, width, height, bg_color=MIDNIGHT_CARD, border_color=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()
        return shape

    def add_arena_header(slide, title, category="YFWD • STRATEGIC PLATFORM BRIEFING 2026"):
        # Header category tag
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title.upper()
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(3)

    def add_footer(slide, current_page, total_pages=15, source="SOURCE: DLD, VALUSTRAT, THE NATIONAL, GULF NEWS 2025–2026"):
        # Footer shape / text
        tx = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.35))
        tf = tx.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"YFWD DUBAI • COM345 STRATEGIC PLAN   |   {source}   |   PAGE {current_page:02d} / {total_pages:02d}"
        p.font.size = Pt(8.5)
        p.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 1: TITLE / COVER (Arena Style)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, MIDNIGHT)

    # Left solid Cobalt Hero Block
    add_card(slide1, Inches(0.8), Inches(1.3), Inches(4.5), Inches(5.3), bg_color=COBALT)
    tb_left = slide1.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(3.9), Inches(4.7))
    tf_l = tb_left.text_frame
    tf_l.word_wrap = True

    p = tf_l.paragraphs[0]
    p.text = "EXECUTIVE MANDATE"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf_l.add_paragraph()
    p2.text = "DUBAI LIFE.\nSMARTER FINANCES."
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(12)

    p3 = tf_l.add_paragraph()
    p3.text = "A regulated behavioral-change and financial literacy framework empowering Dubai's 25–34 generation through the Educate-Equip-Connect methodology."
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = WHITE
    p3.space_before = Pt(14)

    p4 = tf_l.add_paragraph()
    p4.text = "12-MONTH PLAN  •  AED 50,000 TOTAL BUDGET"
    p4.font.size = Pt(10)
    p4.font.bold = True
    p4.font.color.rgb = WHITE
    p4.space_before = Pt(24)

    # Right: 4 Team Bento Blocks
    team_data = [
        ("SECTION 01 • LAREB", "MARKET REALITIES & PRESSURES", "470 daily newcomers, +6% rent inflation forecast, AED 6 Salik peak tolls, and AECB BNPL credit reporting integration."),
        ("SECTION 02 • GLEB", "AUDIENCE & REALLOCATION MODEL", "Dubai residents 25–34, persona Rayan, Educate-Equip-Connect strategy, and unlocking AED 460/mo (+92% savings)."),
        ("SECTION 03 • AMALIYA", "TACTICS, CHANNELS & BUDGET", "2x weekly reels, 6 interactive workshops, free web portal, 30-Day challenge, Zywa/Fazaa deals, and AED 50k budget."),
        ("SECTION 04 • AKBAR", "SMART TARGETS & CONTINUOUS AUDIT", "150K impressions (4% CTR), 5K tool downloads, 1K challenge participants, and 3-phase formative/process/summative review.")
    ]

    for i, (tag, title, desc) in enumerate(team_data):
        rx = Inches(5.6) + (i % 2) * Inches(3.55)
        ry = Inches(1.3) + (i // 2) * Inches(2.75)
        add_card(slide1, rx, ry, Inches(3.35), Inches(2.55), bg_color=MIDNIGHT_CARD, border_color=COBALT if i == 0 else None)
        
        tb = slide1.shapes.add_textbox(rx + Inches(0.25), ry + Inches(0.2), Inches(2.85), Inches(2.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(6)

    add_footer(slide1, 1)

    slide1.notes_slide.notes_text_frame.text = (
        "Welcome everyone. Today we present 'Dubai Life. Smarter Finances.'—a 12-month strategic communication plan "
        "designed for the Youth Financial Well-being Dubai (YFWD). Dubai is thriving, with hundreds moving in daily, "
        "yet young residents aged 25–34 face severe cost-of-living pressures. Our mission is to bridge the gap between "
        "rising costs and existing institutional support using our proven Educate-Equip-Connect model with a lean AED 50,000 budget."
    )

    # ==========================================
    # SLIDE 2: EXECUTIVE SUMMARY & BLUEPRINT
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, MIDNIGHT)
    add_arena_header(slide2, "ONE PLATFORM. THREE PERSPECTIVES.", "EXECUTIVE OVERVIEW")

    exec_cards = [
        ("01", "THE CHALLENGE", "MACRO SQUEEZE", "470 daily newcomers outpace housing supply. Rents climbing 6% with flat wages, Salik tolls +50%, and everyday BNPL installments registered on UAE credit bureau files.", MIDNIGHT_CARD, COBALT),
        ("02", "THE AUDIENCE", "YOUTH 25–34", "Young professionals, fresh university graduates, and private tenants navigating independent living expenses, lifestyle inflation, and social media peer pressure.", MIDNIGHT_CARD, WHITE),
        ("03", "THE STRATEGY", "TRIAD ARCHITECTURE", "Demystifying institutional tools (Educate), providing interactive web planners & 30-Day challenge (Equip), and unlocking Zywa & Fazaa discounts (Connect).", COBALT, WHITE),
        ("04", "THE IMPACT", "AED 50K CAPITAL ROI", "150K impressions, 5,000 tool downloads, 1,000 active budget savers, 300 workshop attendees, and 4 institutional commercial partnerships in 12 months.", MIDNIGHT_CARD, COBALT),
    ]

    card_w = Inches(2.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (num, title, tag, body, bg_col, num_col) in enumerate(exec_cards):
        x = start_x + i * (card_w + gap)
        add_card(slide2, x, Inches(1.5), card_w, Inches(5.1), bg_color=bg_col)
        tb = slide2.shapes.add_textbox(x + Inches(0.25), Inches(1.7), card_w - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = num_col

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = tag
        p3.font.size = Pt(9.5)
        p3.font.bold = True
        p3.font.color.rgb = WHITE if bg_col == COBALT else COBALT
        p3.space_before = Pt(2)

        p4 = tf.add_paragraph()
        p4.text = body
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = WHITE if bg_col == COBALT else TEXT_BODY
        p4.space_before = Pt(12)

    add_footer(slide2, 2)

    # ==========================================
    # SLIDE 3: MARKET REALITIES (Arena Multi-Column Banner)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, MIDNIGHT)
    add_arena_header(slide3, "A MARKET SQUEEZED ACROSS ALL KEY METRICS", "LAREB'S PART • SECTION 1: BACKGROUND")

    metrics_4 = [
        ("DEMOGRAPHIC PRESSURE", "470", "Daily Influx (The National, 2025)", "Dubai's population is expanding faster than housing completions, driving sustained rental demand and occupancy pressure in tenant districts.", COBALT),
        ("HOUSING VS WAGE GAP", "+6%", "2026 Rent Rise vs 0% Salary Growth", "Residential rents projected up by up to 6% in 2026 (Abbas, 2025), while average salaries saw 0% increase in 2025 (ValuStrat, 2025).", WHITE),
        ("COMMUTING OVERHEAD", "AED 6", "Peak Salik Toll (+50% Increase)", "Peak-hour Salik gates increased from AED 4 to AED 6 (Gulf Business, 2024), adding up to AED 480–520 monthly in unavoidable commuting tolls.", COBALT),
        ("HIDDEN CREDIT RISK", "AECB", "BNPL Installments on Credit Bureau", "Tabby & Tamara BNPL purchases are now reported to the Etihad Credit Bureau (Gulf News, 2026), directly damaging future loan and rental approvals.", ROSE)
    ]

    for i, (tag, stat, subtitle, desc, col) in enumerate(metrics_4):
        x = start_x + i * (card_w + gap)
        add_card(slide3, x, Inches(1.5), card_w, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide3.shapes.add_textbox(x + Inches(0.25), Inches(1.7), card_w - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = stat
        p2.font.size = Pt(36)
        p2.font.bold = True
        p2.font.color.rgb = col
        p2.space_before = Pt(6)

        p3 = tf.add_paragraph()
        p3.text = subtitle
        p3.font.size = Pt(10)
        p3.font.bold = True
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(2)

        p4 = tf.add_paragraph()
        p4.text = desc
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = TEXT_BODY
        p4.space_before = Pt(12)

    add_footer(slide3, 3)

    # ==========================================
    # SLIDE 4: PROBLEM VS OPPORTUNITY (Split Screen)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, MIDNIGHT)
    add_arena_header(slide4, "THE INFORMATION PARADOX: PROBLEM VS. OPPORTUNITY", "LAREB'S PART • SECTION 2: PROBLEMS & OPPORTUNITIES")

    # Left: Midnight Block
    add_card(slide4, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=MIDNIGHT_CARD)
    tb_p = slide4.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(4.7))
    tf_p = tb_p.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "THE COMMUNICATION PROBLEM"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ROSE

    p2 = tf_p.add_paragraph()
    p2.text = "LOW INSTITUTIONAL AWARENESS & VIRAL MISINFORMATION"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)

    prob_bullets = [
        "Information Asymmetry: Official tools like the DLD Rental Index and AECB credit reports exist, but young adults do not know they exist.",
        "Unregulated Social Media Flex: Youth rely on TikTok/IG where unverified influencers promote get-rich-quick schemes, prompting UAE authorities to license finance influencers (Issa, 2025).",
        "Unconscious Spending Drain: Daily food deliveries, peak Salik tolls, and unmonitored subscriptions accumulate silently into hundreds of lost dirhams monthly."
    ]
    for b in prob_bullets:
        p_b = tf_p.add_paragraph()
        p_b.text = f"✕  {b}"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = TEXT_BODY
        p_b.space_before = Pt(10)

    # Right: Solid Cobalt Block
    add_card(slide4, Inches(6.833), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=COBALT)
    tb_o = slide4.shapes.add_textbox(Inches(7.133), Inches(1.7), Inches(5.1), Inches(4.7))
    tf_o = tb_o.text_frame
    tf_o.word_wrap = True

    p = tf_o.paragraphs[0]
    p.text = "THE STRATEGIC OPPORTUNITY"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf_o.add_paragraph()
    p2.text = "YFWD AS THE TRUSTED PUBLIC ADVOCATE"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)

    opp_bullets = [
        "Early Preventive Public Mandate: Intervene proactively before living cost pressures become unmanageable social strain.",
        "Meet Youth on Native Digital Channels: Deliver verified, accredited micro-learning directly on TikTok and Instagram using licensed UAE financial creators.",
        "Demystify Official UAE Infrastructure: Translate DLD rental caps, tenant rights, and credit scoring into accessible web calculators and survival toolkits."
    ]
    for b in opp_bullets:
        p_b = tf_o.add_paragraph()
        p_b.text = f"✓  {b}"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = WHITE
        p_b.space_before = Pt(10)

    add_footer(slide4, 4)

    # ==========================================
    # SLIDE 5: TARGET PUBLICS & PERSONA PROFILE (Gleb)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, MIDNIGHT)
    add_arena_header(slide5, "TARGET PUBLICS & PERSONA ARCHITECTURE", "GLEB'S PART • AUDIENCE ARCHITECTURE")

    # Left: Audience Segments
    add_card(slide5, Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.1), bg_color=MIDNIGHT_CARD)
    tb_aud = slide5.shapes.add_textbox(Inches(1.05), Inches(1.7), Inches(4.0), Inches(4.7))
    tf_aud = tb_aud.text_frame
    tf_aud.word_wrap = True

    p = tf_aud.paragraphs[0]
    p.text = "AUDIENCE SEGMENTATION"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COBALT

    p2 = tf_aud.add_paragraph()
    p2.text = "PRIMARY PUBLIC (25–34)"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(6)

    p3 = tf_aud.add_paragraph()
    p3.text = "• Early career corporate & tech employees\n• Recent university graduates entering jobs\n• Private apartment tenants facing renewals\n• Heavy consumers of vertical social video"
    p3.font.size = Pt(10.5)
    p3.font.color.rgb = TEXT_BODY
    p3.space_before = Pt(4)

    p4 = tf_aud.add_paragraph()
    p4.text = "SECONDARY PUBLICS (ENABLERS)"
    p4.font.size = Pt(13)
    p4.font.bold = True
    p4.font.color.rgb = WHITE
    p4.space_before = Pt(16)

    p5 = tf_aud.add_paragraph()
    p5.text = "• Universities & colleges (graduate pipelines)\n• Corporate employers & HR wellness divisions\n• Commercial banks, Zywa fintech & Fazaa"
    p5.font.size = Pt(10.5)
    p5.font.color.rgb = TEXT_BODY
    p5.space_before = Pt(4)

    # Right: Persona Rayan
    add_card(slide5, Inches(5.6), Inches(1.5), Inches(6.933), Inches(5.1), bg_color=MIDNIGHT_CARD)
    tb_ray = slide5.shapes.add_textbox(Inches(5.9), Inches(1.7), Inches(6.333), Inches(4.7))
    tf_ray = tb_ray.text_frame
    tf_ray.word_wrap = True

    p = tf_ray.paragraphs[0]
    p.text = "REPRESENTATIVE PERSONA PROFILE"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COBALT

    p2 = tf_ray.add_paragraph()
    p2.text = "RAYAN, 27 • MARKETING ASSOCIATE (BUSINESS BAY)"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)

    p3 = tf_ray.add_paragraph()
    p3.text = "MONTHLY NET SALARY: AED 6,000 / MO"
    p3.font.size = Pt(12)
    p3.font.bold = True
    p3.font.color.rgb = COBALT
    p3.space_before = Pt(2)

    ray_details = [
        "Current Pressures: Rent takes 45% (AED 2,700). Food deliveries take AED 900. Salik takes AED 150. Active Tabby installments.",
        "Savings Reality: Saves only AED 500/mo (8.3%). Zero emergency cushion for lease renewals. Anxious about rent increases.",
        "Campaign Fit: Responds to bite-sized video hacks, free Excel budget planners, and Fazaa/Zywa verified merchant discounts."
    ]
    for d in ray_details:
        p_d = tf_ray.add_paragraph()
        p_d.text = f"›  {d}"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = TEXT_BODY
        p_d.space_before = Pt(8)

    p_quote = tf_ray.add_paragraph()
    p_quote.text = '"I make good money, but at the end of the month almost nothing is left. Rents are going up, deliveries add up, and I don\'t know where to start saving without sacrificing my life in Dubai."'
    p_quote.font.size = Pt(10.5)
    p_quote.font.italic = True
    p_quote.font.color.rgb = TEXT_MUTED
    p_quote.space_before = Pt(14)

    add_footer(slide5, 5)

    # ==========================================
    # SLIDE 6: CORE STRATEGY (Triad Architecture)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, MIDNIGHT)
    add_arena_header(slide6, "THE STRATEGIC TRIAD: EDUCATE • EQUIP • CONNECT", "GLEB'S PART • STRATEGY & MESSAGING")

    # Banner
    add_card(slide6, Inches(0.8), Inches(1.5), Inches(11.733), Inches(1.1), bg_color=COBALT)
    tb_m = slide6.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(0.9))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = 'CAMPAIGN SLOGAN: "DUBAI LIFE. SMARTER FINANCES."'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p2 = tf_m.add_paragraph()
    p2.text = "Helping young people manage Dubai's cost of living through informed financial decisions, practical planning, and access to real cost-saving opportunities."
    p2.font.size = Pt(11)
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(2)

    # 3 Triad Bento Cards
    triad = [
        ("01", "EDUCATE", "COGNITIVE LITERACY", "• 2x weekly TikTok & IG reels on rent caps & Salik routes\n• 6 practical workshops in universities & employers\n• Partnering with licensed UAE financial creators\n\nTarget: 150K Impressions • 4% Engagement Rate"),
        ("02", "EQUIP", "TOOLS FOR ACTION", "• Free campaign web portal with custom Dubai Budget Planner\n• Interactive Savings Calculator for expense shifting\n• Gamified 30-Day Budget Challenge with weekly sprints\n\nTarget: 5,000 Downloads • 1,000 Challenge Participants"),
        ("03", "CONNECT", "INSTITUTIONAL DEALS", "• Strategic fintech partnership with Zywa prepaid card\n• Fazaa platform integration for verified merchant deals\n• DLD rental index promotion and zero-fee bank accounts\n\nTarget: 4 Signed Alliances • 300 Attendees")
    ]

    w3 = Inches(3.72)
    gap3 = Inches(0.28)
    for i, (num, title, tag, body) in enumerate(triad):
        x = start_x + i * (w3 + gap3)
        add_card(slide6, x, Inches(2.8), w3, Inches(3.8), bg_color=MIDNIGHT_CARD)
        tb = slide6.shapes.add_textbox(x + Inches(0.25), Inches(3.0), w3 - Inches(0.5), Inches(3.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = tag
        p3.font.size = Pt(9.5)
        p3.font.bold = True
        p3.font.color.rgb = COBALT
        p3.space_before = Pt(2)

        p4 = tf.add_paragraph()
        p4.text = body
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = TEXT_BODY
        p4.space_before = Pt(8)

    add_footer(slide6, 6)

    # ==========================================
    # SLIDE 7: SAVINGS PROOF OF CONCEPT (Gleb)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, MIDNIGHT)
    add_arena_header(slide7, "THE AED 6,000 SIMULATION: +92% SAVINGS SURGE", "GLEB'S PART • INTERACTIVE PROOF OF CONCEPT")

    # Left: Explanation Card
    add_card(slide7, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=MIDNIGHT_CARD)
    tb_sc = slide7.shapes.add_textbox(Inches(1.1), Inches(1.7), Inches(5.1), Inches(4.7))
    tf_sc = tb_sc.text_frame
    tf_sc.word_wrap = True

    p = tf_sc.paragraphs[0]
    p.text = "REALLOCATING AED 460 / MONTH"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COBALT

    p2 = tf_sc.add_paragraph()
    p2.text = "IDENTICAL INCOME. NEARLY DOUBLE THE SAVINGS."
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_before = Pt(4)

    p3 = tf_sc.add_paragraph()
    p3.text = (
        "Within an unchanged monthly income of AED 6,000, simple lifestyle optimizations "
        "reclaim AED 460 in leakage without degrading living standards:\n\n"
        "• Meal Prep & Groceries: -AED 180 (cut 3 takeaway deliveries/wk)\n"
        "• Off-Peak Transport & Metro: -AED 120 (avoid peak Salik AED 6 gates)\n"
        "• Subscription Audit: -AED 80 (cancel 2 dormant streaming/app subs)\n"
        "• Partner Discounts (Fazaa/Zywa): -AED 80 (dining & retail savings)\n\n"
        "Result: Monthly savings surge from AED 500 to AED 960 (+92% increase!). "
        "Over 12 months, this accumulates to an extra AED 5,520 in safety buffer."
    )
    p3.font.size = Pt(10.5)
    p3.font.color.rgb = TEXT_BODY
    p3.space_before = Pt(6)

    # Right: Data Table Card
    add_card(slide7, Inches(6.833), Inches(1.5), Inches(5.7), Inches(5.1), bg_color=MIDNIGHT_CARD)
    
    rows, cols = 7, 4
    table_shape = slide7.shapes.add_table(rows, cols, Inches(7.0), Inches(1.7), Inches(5.366), Inches(4.5))
    table = table_shape.table
    table.columns[0].width = Inches(1.9)
    table.columns[1].width = Inches(1.1)
    table.columns[2].width = Inches(1.1)
    table.columns[3].width = Inches(1.266)

    headers = ["Category", "Before", "After", "Monthly Delta"]
    data = [
        ["Rent & Utilities", "AED 2,700", "AED 2,700", "AED 0 (Fixed)"],
        ["Food & Groceries", "AED 1,300", "AED 1,120", "-AED 180"],
        ["Transport & Salik", "AED 800", "AED 680", "-AED 120"],
        ["Lifestyle & Subs", "AED 700", "AED 540", "-AED 160"],
        ["Net Monthly Savings", "AED 500", "AED 960", "+AED 460 (+92%)"],
        ["Total Budget", "AED 6,000", "AED 6,000", "Net Balanced"]
    ]

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COBALT
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = WHITE

    for r, row in enumerate(data):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = MIDNIGHT_CARD2 if r % 2 == 0 else MIDNIGHT_CARD
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                if r == 4:
                    p.font.bold = True
                    p.font.color.rgb = COBALT
                else:
                    p.font.color.rgb = TEXT_BODY

    add_footer(slide7, 7)

    # ==========================================
    # SLIDE 8: CHANNELS & TACTICS - EDUCATE (Amaliya)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, MIDNIGHT)
    add_arena_header(slide8, "PILLAR 1: EDUCATE — CONTENT & WORKSHOPS", "AMALIYA'S PART • TACTICS IN ACTION")

    cards_e = [
        ("DIGITAL CAMPAIGN", "2X WEEKLY VIDEO SERIES", "2x weekly short-form videos covering:\n• How to check the DLD Rental Index before lease signing.\n• Decoding your Salik bills & using public transit hacks.\n• BNPL hidden penalties & protecting your AECB credit score.\n• Smart supermarket grocery shopping in Dubai.\n\nBudget: AED 9,000 prod. + AED 10,000 paid ads."),
        ("CAMPUS & WORKPLACE ROADSHOWS", "SIX IN-PERSON WORKSHOPS", "Delivered in partnership with UAE universities & employers:\n• Free venue partnerships to eliminate facility costs.\n• Interactive roleplay: Negotiating rent renewal with landlords.\n• Guest financial advisors vetted under UAE licensing laws.\n• Target: 300+ total in-person attendees across 12 months.\n\nBudget: AED 9,000 (AED 1,500/workshop)."),
        ("AUTHORITY & TRUST", "LICENSED FIN-CREATORS", "Engaging 4 micro-influencers licensed under the UAE Media Council:\n• Authentic 'My First Year in Dubai' spending breakdowns.\n• Real receipts challenge: Tracking daily coffees vs metro passes.\n• Directing followers to the YFWD free budget webpage tools.\n\nBudget: AED 4,000 (AED 1,000/creator)."),
    ]

    for i, (tag, title, body) in enumerate(cards_e):
        x = start_x + i * (w3 + gap3)
        add_card(slide8, x, Inches(1.5), w3, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide8.shapes.add_textbox(x + Inches(0.25), Inches(1.7), w3 - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(10)

    add_footer(slide8, 8)

    # ==========================================
    # SLIDE 9: CHANNELS & TACTICS - EQUIP (Amaliya)
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, MIDNIGHT)
    add_arena_header(slide9, "PILLAR 2: EQUIP — DIGITAL TOOLS & CHALLENGE", "AMALIYA'S PART • TACTICS IN ACTION")

    cards_eq = [
        ("DIGITAL TOOLKIT", "CAMPAIGN WEB PORTAL", "• Dedicated, mobile-optimized webpage.\n• Free downloadable Excel/Notion Dubai Budget Planner.\n• Interactive Savings Calculator to simulate expense tweaks.\n• Downloadable quick-reference PDF guides for renters & expats.\n\nBudget: AED 6,000 web development."),
        ("HABIT GAMIFICATION", "30-DAY BUDGET CHALLENGE", "Structured 4-week micro-challenge for young professionals:\n• Week 1: Audit & purge dormant subscriptions.\n• Week 2: Salik & Dubai Metro route trial.\n• Week 3: Meal prep & grocery discount week.\n• Week 4: Setting up an emergency cash reserve.\n\nTarget: 1,000 active enrollments."),
        ("OFFLINE TO ONLINE", "CAMPUS & TRANSIT QR POSTERS", "• High-traffic physical placement at university campuses, metro stations, and partner co-working hubs.\n• Scannable QR codes leading directly to free planner tools.\n• Cost-efficient print budget using reusable materials.\n\nBudget: AED 2,000 print deployment."),
    ]

    for i, (tag, title, body) in enumerate(cards_eq):
        x = start_x + i * (w3 + gap3)
        add_card(slide9, x, Inches(1.5), w3, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide9.shapes.add_textbox(x + Inches(0.25), Inches(1.7), w3 - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(10)

    add_footer(slide9, 9)

    # ==========================================
    # SLIDE 10: CHANNELS & TACTICS - CONNECT (Amaliya)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, MIDNIGHT)
    add_arena_header(slide10, "PILLAR 3: CONNECT — FINTECHS & ECOSYSTEM DEALS", "AMALIYA'S PART • TACTICS IN ACTION")

    cards_c = [
        ("FINTECH ALLIANCE", "ZYWA APP & PREPAID CARD", "• Collaboration with Zywa, the UAE youth money app.\n• Teaches real-time transaction tracking & automated expense categorization.\n• Prepaid card model prevents overdraft and enforces disciplined spending habits.\n• Direct educational co-branding inside the app."),
        ("DISCOUNT PLATFORM", "FAZAA NETWORK INTEGRATION", "• Direct partnership with Fazaa discount platform for eligible youth.\n• Verified discounts on daily essentials, supermarket groceries, and entertainment.\n• Transparent eligibility guidelines communicated upfront to build trust.\n• Unlocks AED 80–150 monthly per youth."),
        ("GOVERNMENT & BANKING", "DLD & BANKING ALLIANCES", "• Property portals: Integrating verified DLD rental calculator data into tenant workshops.\n• Commercial banks: Clear advice on no-fee savings accounts and avoiding credit card interest traps.\n• Credit Bureau: Teaching youth how to pull and read AECB credit reports annually."),
    ]

    for i, (tag, title, body) in enumerate(cards_c):
        x = start_x + i * (w3 + gap3)
        add_card(slide10, x, Inches(1.5), w3, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide10.shapes.add_textbox(x + Inches(0.25), Inches(1.7), w3 - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(14)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(10)

    add_footer(slide10, 10)

    # ==========================================
    # SLIDE 11: 12-MONTH BUDGET ALLOCATION (Amaliya)
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11, MIDNIGHT)
    add_arena_header(slide11, "12-MONTH FINANCIAL BUDGET: AED 50,000 ALLOCATION", "AMALIYA'S PART • FINANCIAL ALLOCATION")

    # Left: Budget Table Card
    add_card(slide11, Inches(0.8), Inches(1.5), Inches(6.8), Inches(5.1), bg_color=MIDNIGHT_CARD)
    
    b_rows, b_cols = 11, 3
    b_table_shape = slide11.shapes.add_table(b_rows, b_cols, Inches(1.0), Inches(1.7), Inches(6.4), Inches(4.5))
    b_table = b_table_shape.table
    b_table.columns[0].width = Inches(3.6)
    b_table.columns[1].width = Inches(1.4)
    b_table.columns[2].width = Inches(1.4)

    b_headers = ["Budget Allocation Item", "AED Amount", "% of Total"]
    b_data = [
        ["Social Media Advertising (IG/TikTok Paid)", "AED 10,000", "20.0%"],
        ["Educational Content Creation & Video Prod.", "AED 9,000", "18.0%"],
        ["Six Interactive In-Person Workshops", "AED 9,000", "18.0%"],
        ["Campaign Webpage & Digital Tools", "AED 6,000", "12.0%"],
        ["Research & Evaluation (Baseline + Summative)", "AED 5,000", "10.0%"],
        ["Creator / Influencer Collaborations", "AED 4,000", "8.0%"],
        ["Contingency Fund (Agile reallocation)", "AED 3,000", "6.0%"],
        ["Printing & Campus/Transit QR Posters", "AED 2,000", "4.0%"],
        ["Partner Coordination & Administration", "AED 2,000", "4.0%"],
        ["Total Campaign Budget (12 Months)", "AED 50,000", "100.0%"]
    ]

    for c, h in enumerate(b_headers):
        cell = b_table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COBALT
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9.5)
            p.font.bold = True
            p.font.color.rgb = WHITE

    for r, row in enumerate(b_data):
        for c, val in enumerate(row):
            cell = b_table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = MIDNIGHT_CARD2 if r % 2 == 0 else MIDNIGHT_CARD
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.5)
                if r == 9:
                    p.font.bold = True
                    p.font.color.rgb = COBALT
                else:
                    p.font.color.rgb = TEXT_BODY

    # Right: Cost-Efficiency Card
    add_card(slide11, Inches(7.8), Inches(1.5), Inches(4.733), Inches(5.1), bg_color=COBALT)
    tb_eff = slide11.shapes.add_textbox(Inches(8.05), Inches(1.7), Inches(4.2), Inches(4.7))
    tf_eff = tb_eff.text_frame
    tf_eff.word_wrap = True

    p = tf_eff.paragraphs[0]
    p.text = "BUDGET EFFICIENCY PRINCIPLES"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE

    eff_points = [
        "Free Partner Venues: University and incubator partnerships eliminate AED 15k+ in facility hire fees.",
        "Volunteer Speakers: Sourcing certified financial planners and credit bureau analysts pro bono.",
        "Partner Cross-Promotion: Utilizing partner newsletters (Zywa, universities) for zero-cost reach.",
        "Reusable Digital Assets: Modular infographics and templates reused across digital channels.",
        "Supplier Price Locks: Lock quotation rates before disbursement; monthly audits to prevent creep.",
        "Agile Contingency: AED 3,000 buffer held for Month 6 reallocation toward highest-performing tactic."
    ]

    for pt in eff_points:
        p_pt = tf_eff.add_paragraph()
        p_pt.text = f"✓  {pt}"
        p_pt.font.size = Pt(10)
        p_pt.font.color.rgb = WHITE
        p_pt.space_before = Pt(6)

    add_footer(slide11, 11)

    # ==========================================
    # SLIDE 12: SMART OBJECTIVES & TARGETS (Akbar)
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12, MIDNIGHT)
    add_arena_header(slide12, "THREE SMART OBJECTIVES: TARGETS & BENCHMARKS", "AKBAR'S PART • SECTION 3: OBJECTIVES")

    cards_obj = [
        ("OBJECTIVE 1 • EDUCATE", "150K", "Total Impressions", "4.0% Avg Engagement Rate", "To increase awareness of financial literacy and YFWD support resources among Dubai residents aged 25–34 within 6 months of launch.\n\nAudit: Social analytics, video views, comments & shares."),
        ("OBJECTIVE 2 • EQUIP", "5,000", "Tool Downloads / Uses", "1,000 Challenge Participants", "To drive proactive financial habit formation through digital adoption within 12 months.\n\nAudit: Google Analytics, QR telemetry, and completion forms."),
        ("OBJECTIVE 3 • CONNECT", "300+", "Workshop Attendees", "4 Active Partnerships", "To build sustainable community and commercial connections for long-term impact within 12 months.\n\nAudit: Attendance registers, signed MOUs & redeemed partner perks.")
    ]

    for i, (tag, stat1, label1, label2, body) in enumerate(cards_obj):
        x = start_x + i * (w3 + gap3)
        add_card(slide12, x, Inches(1.5), w3, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide12.shapes.add_textbox(x + Inches(0.25), Inches(1.7), w3 - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = tag
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = COBALT

        p2 = tf.add_paragraph()
        p2.text = stat1
        p2.font.size = Pt(36)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = label1
        p3.font.size = Pt(10)
        p3.font.bold = True
        p3.font.color.rgb = COBALT
        p3.space_before = Pt(2)

        p4 = tf.add_paragraph()
        p4.text = f"KPI Target: {label2}"
        p4.font.size = Pt(11)
        p4.font.bold = True
        p4.font.color.rgb = WHITE
        p4.space_before = Pt(8)

        p5 = tf.add_paragraph()
        p5.text = body
        p5.font.size = Pt(10.5)
        p5.font.color.rgb = TEXT_BODY
        p5.space_before = Pt(8)

    add_footer(slide12, 12)

    # ==========================================
    # SLIDE 13: 3-PHASE EVALUATION (Akbar)
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide13, MIDNIGHT)
    add_arena_header(slide13, "OUR JOURNEY TO IMPACT: 3-PHASE EVALUATION", "AKBAR'S PART • SECTION 10: EVALUATION")

    phases = [
        ("MONTHS 1–2", "01", "FORMATIVE EVALUATION", "• Desktop Research & Social Listening: Monitor hashtags & online discussions around Dubai living costs and BNPL sentiment.\n• Pre-Campaign Digital Survey: Measure baseline awareness of YFWD, DLD rental index, and financial stress.\n\nOutput: Validated baseline benchmark against which Year 1 delta is measured."),
        ("MONTHS 3–8", "02", "PROCESS EVALUATION & AGILITY", "• Monthly Social Analytics: If engagement drops below 4%, immediately pivot formats toward high-performing Reels/creators.\n• Webpage Conversion Tracking: Monitor Google Analytics funnel and challenge completions.\n• Month 6 Agility Audit: Reallocate unused contingency (AED 3k) to top ROI channels."),
        ("MONTHS 10–12", "03", "SUMMATIVE EVALUATION & ROI", "• Month 12 Post-Campaign Survey: Measure shifts in financial literacy, self-reported monthly savings, and trust in YFWD.\n• SMART Objective Audit: Verify actuals against 150k impressions, 5k downloads, 1k participants, 4 partner contracts.\n• Final ROI Report: Reach per AED spent to advise government on long-term policy.")
    ]

    for i, (time_h, num, title, body) in enumerate(phases):
        x = start_x + i * (w3 + gap3)
        add_card(slide13, x, Inches(1.5), w3, Inches(5.1), bg_color=MIDNIGHT_CARD)
        tb = slide13.shapes.add_textbox(x + Inches(0.25), Inches(1.7), w3 - Inches(0.5), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = time_h
        p.font.size = Pt(9.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_MUTED

        p2 = tf.add_paragraph()
        p2.text = num
        p2.font.size = Pt(28)
        p2.font.bold = True
        p2.font.color.rgb = COBALT
        p2.space_before = Pt(4)

        p3 = tf.add_paragraph()
        p3.text = title
        p3.font.size = Pt(13)
        p3.font.bold = True
        p3.font.color.rgb = WHITE
        p3.space_before = Pt(2)

        p4 = tf.add_paragraph()
        p4.text = body
        p4.font.size = Pt(10.5)
        p4.font.color.rgb = TEXT_BODY
        p4.space_before = Pt(10)

    add_footer(slide13, 13)

    # ==========================================
    # SLIDE 14: RISK MANAGEMENT PLAYBOOK
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide14, MIDNIGHT)
    add_arena_header(slide14, "RISK MANAGEMENT & CONTINGENCY PLAYBOOK", "OPERATIONAL RESILIENCE")

    risks = [
        ("RISK 1: SOCIAL ENGAGEMENT BELOW 4%", "MEDIA RISK", "Pivot video style from informational monologues to relatable day-in-the-life skits. Reallocate portion of AED 10k paid ad budget into A/B testing high-performing hooks."),
        ("RISK 2: WORKSHOP ATTENDANCE DROP", "EVENTS RISK", "Incentivize attendance through certified CPD certificates and complimentary partner discount codes. Shift sessions to hybrid live-stream webinars."),
        ("RISK 3: PARTNER ONBOARDING DELAYS", "ALLIANCE RISK", "Prioritize fintechs with pre-existing open APIs (Zywa) and established institutional programs (Fazaa). Offer co-branded thought leadership as mutual value."),
        ("RISK 4: UAE LICENSING COMPLIANCE", "LEGAL RISK", "Strictly pre-vet all content and partner contracts in alignment with UAE National Media Council guidelines and licensed financial influencer rules.")
    ]

    for i, (metric, title, desc) in enumerate(risks):
        x = start_x + (i % 2) * (Inches(5.7) + Inches(0.33))
        y = Inches(1.5) + (i // 2) * (Inches(2.5) + Inches(0.25))
        add_card(slide14, x, y, Inches(5.7), Inches(2.5), bg_color=MIDNIGHT_CARD)
        
        tb = slide14.shapes.add_textbox(x + Inches(0.25), y + Inches(0.2), Inches(5.2), Inches(2.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = metric
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(9.5)
        p2.font.bold = True
        p2.font.color.rgb = COBALT
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_BODY
        p3.space_before = Pt(6)

    add_footer(slide14, 14)

    # ==========================================
    # SLIDE 15: 12-MONTH ROADMAP & CONCLUSION
    # ==========================================
    slide15 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide15, MIDNIGHT)
    add_arena_header(slide15, "12-MONTH ROADMAP: DELIVERING LASTING PROSPERITY", "EXECUTION ROADMAP & SIGN-OFF")

    q_data = [
        ("Q1: MONTHS 1–3", "FOUNDATION & LAUNCH", "• Baseline survey & social listening\n• Launch web portal & budget templates\n• Wave 1 TikTok/IG reels & creator posts\n• Deliver Workshops 1 & 2 (Universities)"),
        ("Q2: MONTHS 4–6", "SCALE & AGILITY REVIEW", "• Launch '30-Day Budget Challenge'\n• Onboard Zywa & Fazaa partnerships\n• Deliver Workshops 3 & 4 (Employers)\n• Month 6 Review & contingency reallocation"),
        ("Q3: MONTHS 7–9", "ECOSYSTEM DEEPENING", "• Deliver Workshops 5 & 6 (Hubs)\n• High-impact creator collaborations\n• QR code poster refreshes across Metro\n• Mid-point survey pulse check"),
        ("Q4: MONTHS 10–12", "SUMMATIVE & POLICY", "• Month 12 summative evaluation survey\n• Audit performance vs 3 SMART targets\n• Compile final ROI executive report\n• Present permanent policy to YFWD")
    ]

    for i, (q_title, sub, body) in enumerate(q_data):
        x = start_x + i * (card_w + gap)
        add_card(slide15, x, Inches(1.5), card_w, Inches(5.1), bg_color=COBALT if i == 1 else MIDNIGHT_CARD)
        tb = slide15.shapes.add_textbox(x + Inches(0.2), Inches(1.7), card_w - Inches(0.4), Inches(4.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = q_title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(10)
        p2.font.bold = True
        p2.font.color.rgb = WHITE if i == 1 else COBALT
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = body
        p3.font.size = Pt(10.5)
        p3.font.color.rgb = WHITE if i == 1 else TEXT_BODY
        p3.space_before = Pt(12)

    add_footer(slide15, 15)

    output_path = "Dubai_Life_Smarter_Finances_COM345.pptx"
    prs.save(output_path)
    print(f"Successfully generated DFM Arena styled {output_path} with {len(prs.slides)} slides!")

if __name__ == "__main__":
    create_presentation()
